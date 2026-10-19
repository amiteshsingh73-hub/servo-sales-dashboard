"""Generates a deck on the HO standard template (cover + styled slides with native, editable tables/charts)."""
from __future__ import annotations
import copy, io
import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt, Emu
from lxml import etree
from . import config as C, ho_tables as H, metrics as M
from .periods import fy_label

BANNER, HEADER, ROW_A, ROW_B = "4472C4", "CFD5EA", "E9EBF5", "CFD5EA"
NAVY, GREY = RGBColor(0x1F, 0x3A, 0x93), RGBColor(0xA6, 0xA6, 0xA6)


# ---------------------------------------------------------------- cell formatting
def fmt(col: str, v) -> str:
    if v is None or (isinstance(v, float) and np.isnan(v)): return "-"
    if isinstance(v, str): return v
    c = col.lower()
    if "%" in c:
        if abs(v) > 999: return "n.m."
        return f"{v:.1f}%" if "cagr" in c else f"{v:.0f}%"
    if c in ("s.no.",): return f"{int(v)}"
    if "ltrs" in c: return f"{v:.0f}"
    return f"{v:,.1f}" if abs(v) < 10 else f"{v:,.0f}"


def _border(cell, color="000000", w=6350):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for old in tcPr.findall(qn(tag)): tcPr.remove(old)
        ln = etree.SubElement(tcPr, qn(tag), w=str(w), cap="flat", cmpd="sng", algn="ctr")
        sf = etree.SubElement(ln, qn("a:solidFill")); etree.SubElement(sf, qn("a:srgbClr"), val=color)
        etree.SubElement(ln, qn("a:prstDash"), val="solid")
    # borders must precede fill in schema order -> move fill (if any) to the end
    for f in tcPr.findall(qn("a:solidFill")): tcPr.remove(f); tcPr.append(f)


def _set(cell, text, size=10, bold=False, color="000000", fill=None, align=PP_ALIGN.CENTER):
    cell.text = ""
    tf = cell.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size, r.font.bold, r.font.name = Pt(size), bold, "Arial"
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = cell.margin_right = Emu(40000); cell.margin_top = cell.margin_bottom = Emu(10000)
    if fill:
        cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor.from_string(fill)
    _border(cell)


# ---------------------------------------------------------------- slide helpers
def _layout(prs):
    for l in prs.slide_layouts:
        if l.name == "Title and Content": return l
    return prs.slide_layouts[0]


def new_slide(prs, title: str):
    s = prs.slides.add_slide(_layout(prs))
    for ph in list(s.placeholders):
        if ph.placeholder_format.idx != 0: ph._element.getparent().remove(ph._element)
    if s.shapes.title is not None:
        s.shapes.title.text = title
    return s


def textbox(slide, text, left, top, width, height, size=11, color="003087", bold=False, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(left, top, width, height); tf = tb.text_frame; tf.word_wrap = True
    for i, line in enumerate(text if isinstance(text, list) else [text]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align
        r = p.add_run(); r.text = line; r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = RGBColor.from_string(color)
    return tb


def add_table(slide, df: pd.DataFrame, banner: str | None = None, top=Inches(1.25), unit="KL", note: str | None = None,
              left=Inches(0.4), width=Inches(12.5), max_h=Inches(5.1)):
    df = df.copy()
    n_r, n_c = len(df), len(df.columns)
    rows = n_r + 1 + (1 if banner else 0)
    size = 11 if n_c <= 8 else 10 if n_c <= 12 else 9
    rh = min(Inches(0.4), int(max_h / rows))
    first_w = Inches(2.3 if n_c > 6 else 3.0)
    gf = slide.shapes.add_table(rows, n_c, left, top, width, rh * rows)
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    for a in ("firstRow", "bandRow"): tblPr.set(a, "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is not None: tblPr.remove(sid)
    tbl.columns[0].width = first_w
    for j in range(1, n_c): tbl.columns[j].width = int((width - first_w) / (n_c - 1)) if n_c > 1 else width
    r0 = 0
    if banner:
        tbl.cell(0, 0).merge(tbl.cell(0, n_c - 1))
        _set(tbl.cell(0, 0), banner, 14, True, "FFFFFF", BANNER)
        for j in range(1, n_c): _border(tbl.cell(0, j))
        r0 = 1
    for j, c in enumerate(df.columns):
        _set(tbl.cell(r0, j), str(c), size, True, "000000", HEADER)
    tot_idx = n_r - 1 if n_r and str(df.iloc[-1, 0]).lower().startswith(("total", "gr total")) else -1
    for i in range(n_r):
        fill = ROW_A if i % 2 == 0 else ROW_B
        for j, c in enumerate(df.columns):
            v = df.iloc[i, j]
            _set(tbl.cell(r0 + 1 + i, j), fmt(c, v), size, bold=(i == tot_idx), fill=fill,
                 align=PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER)
    for r in tbl.rows: r.height = rh
    if unit:
        textbox(slide, f"Figs in {unit}", Inches(11.6), Inches(0.7), Inches(1.5), Inches(0.35), 12, "000000", False, PP_ALIGN.RIGHT)
    if note:
        textbox(slide, note, Inches(0.4), Inches(6.55), Inches(12.3), Inches(0.6), 10)
    return gf


def add_bar_chart(slide, cats, series: dict[str, list], left, top, width, height, title=None, colors=None, labels=True):
    cd = CategoryChartData(); cd.categories = list(cats)
    for k, v in series.items(): cd.add_series(k, [None if (x is None or (isinstance(x, float) and np.isnan(x))) else round(float(x), 1) for x in v])
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, left, top, width, height, cd)
    ch = gf.chart
    ch.has_legend = len(series) > 1
    if ch.has_legend: ch.legend.position = XL_LEGEND_POSITION.TOP; ch.legend.include_in_layout = False; ch.legend.font.size = Pt(11)
    ch.has_title = bool(title)
    if title:
        ch.chart_title.text_frame.text = title; ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    cols = colors or [NAVY, GREY]
    for s, c in zip(ch.plots[0].series, cols): s.format.fill.solid(); s.format.fill.fore_color.rgb = c
    pl = ch.plots[0]; pl.gap_width = 60
    if labels:
        pl.has_data_labels = True; dl = pl.data_labels; dl.font.size = Pt(9); dl.number_format = '#,##0'; dl.number_format_is_linked = False
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
    ch.category_axis.tick_labels.font.size = Pt(10); ch.value_axis.tick_labels.font.size = Pt(10)
    ch.value_axis.has_major_gridlines = False
    return gf


# ---------------------------------------------------------------- slide builders
def s_highlights(prs, ctx):
    s = new_slide(prs, f"Major Highlights {ctx.wp['cy'].label}")
    lines = ["• " + h for h in H.highlights(ctx)]
    tb = textbox(s, lines, Inches(0.6), Inches(1.2), Inches(12.1), Inches(5.6), 16, "000000")
    for p in tb.text_frame.paragraphs: p.space_after = Pt(10)


def s_summary(prs, ctx):
    s = new_slide(prs, f"TARGET VS ACHIEVEMENT: {fy_label(ctx.period.fy, True)}")
    t = H.summary_table(ctx)
    add_table(s, t, unit="KL unless stated", note="Primary = channel (stockist) billing. Secondary = stockist to counter sales (ClearBlue excluded; ServoXpert counted in Bazaar).", top=Inches(1.0), max_h=Inches(5.4))


def _tbl(title, fn, banner, note, unit="KL"):
    def b(prs, ctx):
        s = new_slide(prs, title); add_table(s, fn(ctx), banner, unit=unit, note=note)
    return b


def s_trend(prs, ctx):
    s = new_slide(prs, f"Monthly sales trend {fy_label(ctx.period.fy)} vs LY")
    mons = ["Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec", "Jan", "Feb", "Mar"]
    for i, (ttl, df) in enumerate((("Primary sales (KL)", ctx.pri), ("Secondary sales (KL)", ctx.sec))):
        cy, ly = M.monthly(df, ctx.period.fy), M.monthly(df, ctx.period.fy - 1)
        add_bar_chart(s, mons, {fy_label(ctx.period.fy, True): cy.values, fy_label(ctx.period.fy - 1, True): ly.values},
                      Inches(0.3 + 6.4 * i), Inches(1.1), Inches(6.3), Inches(5.3), ttl, labels=False)
    textbox(s, "Source: Primary Sales & Secondary Sales files.", Inches(0.4), Inches(6.55), Inches(8), Inches(0.4), 10)


def s_channel(prs, ctx):
    s = new_slide(prs, "Secondary sales by channel")
    ch = M.compare(ctx.sec, "channel", ctx.ws, order=C.CHANNEL_ORDER)
    add_bar_chart(s, ch.index, {f"CY {ctx.ws['cy'].label}": ch.cy.values, f"LY {ctx.ws['ly'].label}": ch.ly.values},
                  Inches(0.3), Inches(1.1), Inches(6.5), Inches(5.3), "Secondary sales by channel (KL)")
    t = pd.DataFrame({"Channel": ch.index, "CY": ch.cy.values, "LY": ch.ly.values, "% Growth": ch.growth.values, "% Share": ch.share.values})
    add_table(s, t, None, top=Inches(1.6), left=Inches(7.0), width=Inches(6.0), max_h=Inches(3.0))


def s_pri_sec(prs, ctx):
    s = new_slide(prs, "Primary vs Secondary - stockist wise")
    p = ctx.pri[(ctx.pri.date >= ctx.wp["cy"].start) & (ctx.pri.date <= ctx.wp["cy"].end)].groupby("stockist").kl.sum()
    q = ctx.sec[(ctx.sec.date >= ctx.ws["cy"].start) & (ctx.sec.date <= ctx.ws["cy"].end)].groupby("stockist").kl.sum()
    t = pd.DataFrame({"Primary": p, "Secondary": q}).fillna(0)
    t["Pri/Sec ratio"] = t.Primary / t.Secondary.replace(0, np.nan); t["Stock build (KL)"] = t.Primary - t.Secondary
    t = t.sort_values("Primary", ascending=False)
    t.loc["Total"] = [t.Primary.sum(), t.Secondary.sum(), t.Primary.sum() / max(t.Secondary.sum(), 1e-9), t["Stock build (KL)"].sum()]
    t = t.reset_index().rename(columns={"index": "Stockist", "stockist": "Stockist"})
    add_table(s, t, f"{ctx.wp['cy'].label}", note="Pri/Sec ratio > 1 = stockists bought more than they sold (stock build-up).")


def s_pivot(prs, ctx):
    if not ctx.pivot: return
    title, t = ctx.pivot
    t = t.reset_index()
    t.columns = [" | ".join(str(x) for x in c if str(x)) if isinstance(c, tuple) else str(c) for c in t.columns]
    note = None
    if len(t) > 16: t, note = t.head(16), f"Showing first 16 rows of the pivot."
    if t.shape[1] > 13: t, note = t.iloc[:, :13], (note or "") + " Showing first 12 value columns."
    s = new_slide(prs, title or "Custom pivot")
    add_table(s, t, None, unit=None, note=note)


SLIDES = {
    "highlights": ("Major highlights (auto text)", s_highlights),
    "summary": ("Target vs Achievement summary", s_summary),
    "segment": ("Primary - Product segment (Auto / Industrial / Co-branded)",
                _tbl("Primary Sales Performance- Product segment", H.segment_table, "Primary Sales in KL", "Channel sales only.")),
    "stockist": ("Primary - Stockist wise", _tbl("Primary Sales Performance- Stockist wise", H.stockist_table, "Primary Sales in KL",
                                                  "Channel sales only. Balance = FY target – sales so far.")),
    "premium": ("Primary - Premium grade by stockist", _tbl("Gold/ Premium Grade sales plan", H.premium_stockist_table, "Premium grade sales in KL", "Premium = 'Premium' category in material master.")),
    "pgroup": ("Primary - Product group", _tbl("Product Segment sales performance and plan (Channel)", H.product_group_table, "Primary Sales in KL",
                                                "Targets apportioned on LY mix. Speciality = brake fluid, coolant & transmission fluid.")),
    "pack": ("Primary - Pack wise", _tbl("Pack wise sales performance and plan (Channel)", H.pack_table, "Primary Sales in KL", "Targets apportioned on LY mix.")),
    "category": ("Primary - Product category (Premium/Aspirational/…)", _tbl("Primary sales - Product category", H.category_table, "Primary Sales in KL", "Category as per material master.")),
    "sec_channel": ("Secondary - Channel", _tbl("Secondary Sales Performance – Channel", H.sec_channel_table, "Secondary Sales in KL", "RO = IOC retail outlets; Bazaar includes ServoXpert.")),
    "sec_stockist": ("Secondary - Stockist wise", _tbl("Secondary Sales Performance – Stockist wise", H.sec_stockist_table, "Secondary Sales in KL", "")),
    "sec_ro": ("Secondary - RO channel by stockist", _tbl("Secondary Sales Performance – RO Channel", lambda c: H.sec_channel_stockist_table(c, "RO"), "Secondary Sales in KL", "IOC RO Channel only.")),
    "sec_bazaar": ("Secondary - Bazaar by stockist", _tbl("Secondary Sales Performance – Bazaar Channel", lambda c: H.sec_channel_stockist_table(c, "Bazaar"), "Secondary Sales in KL", "")),
    "sec_oem": ("Secondary - OEM by stockist", _tbl("Secondary Sales Performance – OEM Channel", lambda c: H.sec_channel_stockist_table(c, "OEM"), "Secondary Sales in KL", "")),
    "pivot": ("Custom pivot (saved from the Pivot tab)", s_pivot),
    "trend": ("Monthly trend charts (Primary & Secondary vs LY)", s_trend),
    "channel_chart": ("Secondary channel chart + table", s_channel),
    "pri_sec": ("Primary vs Secondary by stockist", s_pri_sec),
}
DEFAULT_SLIDES = ["highlights", "summary", "stockist", "segment", "pgroup", "pack", "premium", "sec_channel", "sec_stockist", "sec_ro", "trend"]


def _drop_slides(prs, keep_first=1):
    ids = prs.slides._sldIdLst
    for sldId in list(ids)[keep_first:]:
        prs.part.drop_rel(sldId.rId); ids.remove(sldId)


def build_deck(ctx: H.Ctx, keys: list[str], template_path=None, cover=True) -> bytes:
    prs = Presentation(str(template_path)) if template_path else Presentation()
    if template_path:
        _drop_slides(prs, 1 if cover else 0)
    else:
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for k in keys:
        SLIDES[k][1](prs, ctx)
    buf = io.BytesIO(); prs.save(buf); return buf.getvalue()
