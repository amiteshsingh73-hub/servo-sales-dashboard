"""LRS Field Sales Analysis format (one sheet per stockist / SSR, one month at a time).

Layout follows 'LRS Field Sales Analysis Format.xlsx':
  * 2 - Sales Analysis      : stockist, month row (Potential / Target / Sales CY / Sales LY / +- / Remarks), cumulative row
  LUBES RESELLER GROUP      : secondary sales by outlet type for a period - total no., cum. no. supplied, CY, LY, % of total

Sources: sales from the Primary and Secondary sales files only (top section = primary, outlet table = secondary;
Total No. = outlets billed in the 12 months to the period end), targets from the Target workbook, and an optional
'SSR Potential.xlsx' in the data folder (annual potential in KL per stockist). Without that file the potential is
derived from market share: last full FY sales of the stockist / assumed IOCL market share (default 11%, range 10-12%).
"""
from __future__ import annotations
import io, re, zipfile
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from . import config as C, metrics as M
from .periods import make_period, windows, mi_of, mi_name, fy_label

ROWS = ["Retail Outlets - Non KSK RO", "Retail Outlets - KSK RO", "Retail Outlets - Nayara Energy", "Bazaar Shop", "OE Dealers", "Fleet"]
POTENTIAL_FILE = "SSR Potential.xlsx"


# ------------------------------------------------------------------ outlet type of every secondary line
def outlet_type(sec: pd.DataFrame) -> pd.Series:
    sub = sec.cust_subgroup.astype(str).str.upper().str.replace(" ", "")
    grp = sec.cust_group.astype(str).str.upper()
    return pd.Series(np.select(
        [sub.eq("NAYARARO") | grp.eq("NAYARA RO"),
         grp.eq("DEALER") & sub.eq("IOC-KSK"),
         grp.eq("DEALER"),
         sec.channel.eq("Bazaar"), sec.channel.eq("OEM"), sec.channel.eq("Fleet")],
        [ROWS[2], ROWS[1], ROWS[0], ROWS[3], ROWS[4], ROWS[5]], "Other"), index=sec.index)


def _cid(x) -> str:
    s = str(x).strip()
    s = s[:-2] if s.endswith(".0") else s
    return s.lstrip("0") or s


# ------------------------------------------------------------------ potential (editable workbook)
def potential_path(data_dir: Path | None = None) -> Path:
    return Path(data_dir or C.DATA_DIR) / POTENTIAL_FILE


def load_potential(data_dir: Path | None = None) -> dict:
    """{(stockist_code, fy): annual potential KL}; fy=None means 'every year'."""
    f = potential_path(data_dir)
    if not f.exists(): return {}
    try:
        df = pd.read_excel(f, engine="calamine")
        cols = {re.sub(r"[^a-z]", "", str(c).lower()): c for c in df.columns}
        code, pot = cols.get("stockistcode"), next((v for k, v in cols.items() if "potential" in k), None)
        fyc = next((v for k, v in cols.items() if k in ("fy", "financialyear")), None)
        out = {}
        for _, r in df.iterrows():
            c, p = pd.to_numeric(r.get(code), errors="coerce"), pd.to_numeric(r.get(pot), errors="coerce")
            if pd.notna(c) and pd.notna(p):
                fy = None
                if fyc is not None and pd.notna(r.get(fyc)):
                    mt = re.findall(r"\d{2,4}", str(r.get(fyc)))
                    if mt: fy = 2000 + int(mt[-1][-2:])
                out[(int(c), fy)] = float(p)
        return out
    except Exception:
        return {}


def write_potential_template(stockists: dict, fy: int, data_dir: Path | None = None) -> Path:
    f = potential_path(data_dir)
    pd.DataFrame({"Stockist Code": list(stockists), "Stockist Name": list(stockists.values()), "FY": fy_label(fy),
                  "Annual Potential (KL)": [None] * len(stockists)}).to_excel(f, index=False)
    return f


# ------------------------------------------------------------------ numbers for one stockist / month
def months_available(pri: pd.DataFrame, sec: pd.DataFrame) -> list[int]:
    mis = set(pri.mi.unique()) | set(sec.mi.unique())
    return sorted((int(m) for m in mis), reverse=True)


def _fy_fm(mi: int) -> tuple[int, int]:
    y, m = divmod(mi, 12); m += 1
    fy = y + 1 if m >= C.FY_START_MONTH else y
    return fy, (m - C.FY_START_MONTH) % 12 + 1


def _remark(cy, tgt, ly) -> str:
    parts = []
    if tgt == tgt and tgt: parts.append(f"{cy / tgt * 100:.0f}% of target ({cy - tgt:+,.1f} KL)")
    if ly: parts.append(f"{(cy / ly - 1) * 100:+.0f}% vs LY")
    elif cy: parts.append("no sales LY")
    return "; ".join(parts)


def build(pri, sec, targets, pri_last, sec_last, mi: int, basis="Primary", sec_period="Cumulative",
          potential: dict | None = None, share: float = 0.11, share_range=(0.10, 0.12)) -> list[dict]:
    fy, fm = _fy_fm(mi)
    potential = potential or {}
    base, last = (pri, pri_last) if basis == "Primary" else (sec, sec_last)
    wm = windows(make_period("Month", fy, fm=fm), last)
    wc = windows(make_period("YTD (Apr to month)", fy, fm=fm), last)
    ws_sec = windows(make_period("Month", fy, fm=fm) if sec_period == "Month" else make_period("YTD (Apr to month)", fy, fm=fm), sec_last)
    names = pd.concat([pri[["stockist_code", "stockist"]], sec[["stockist_code", "stockist"]]]).drop_duplicates("stockist_code").set_index("stockist_code").stockist
    sec_t = sec.assign(otype=outlet_type(sec), cid=sec.customer_id.map(_cid))
    year_ago = ws_sec["cy"].end - pd.DateOffset(months=12)
    out = []
    for code in sorted(names.index, key=lambda c: str(names[c])):
        b = base[base.stockist_code == code]
        pot_y = potential.get((int(code), fy), potential.get((int(code), None), np.nan))
        base_ly = float(b[b.fy == fy - 1].kl.sum())          # last full financial year, same basis as the top section
        pot_src = "file" if pot_y == pot_y else ("share" if base_ly > 0 else "none")
        if pot_src == "share": pot_y = base_ly / share
        rng_y = (base_ly / share_range[1], base_ly / share_range[0]) if base_ly > 0 else (np.nan, np.nan)
        r = {"code": int(code), "stockist": names[code], "fy": fy, "fm": fm, "mi": mi, "basis": basis,
             "month_label": mi_name(mi), "cum_label": f"{mi_name(mi_of(fy, 1))} to {mi_name(mi)}",
             "partial": bool(wm["cy"].partial), "data_till": last}
        for key, w, fms in (("month", wm, [fm]), ("cum", wc, list(range(1, fm + 1)))):
            cy, ly = M.kl(b, w["cy"]), M.kl(b, w["ly"])
            tgt = M.target_for(targets, [code], fy, fms, basis)
            f_ = len(fms) / 12
            pot = pot_y * f_ if pot_y == pot_y else np.nan
            r[key] = dict(potential=pot, pot_lo=rng_y[0] * f_, pot_hi=rng_y[1] * f_, target=tgt, cy=cy, ly=ly, diff=cy - ly,
                          remark=_remark(cy, tgt, ly) + (f"; market share {cy / pot * 100:.1f}% of potential" if pot == pot and pot else ""))
        r.update(pot_src=pot_src, pot_base=base_ly, share=share, share_range=share_range, pot_fy=fy_label(fy - 1))
        # ---- secondary by outlet type
        s = sec_t[sec_t.stockist_code == code]
        cy_s, ly_s = M.sel(s, ws_sec["cy"]), M.sel(s, ws_sec["ly"])
        billed_12m = s[(s.date > year_ago) & (s.date <= ws_sec["cy"].end) & (s.kl > 0)]
        rows = []
        for t in ROWS:
            sup = set(cy_s[(cy_s.otype == t) & (cy_s.kl > 0)].cid)
            universe = set(billed_12m[billed_12m.otype == t].cid)
            rows.append({"Channel": t, "Total No.": len(universe | sup), "Cum. No. Supplied": len(sup),
                         "Sales CY": float(cy_s[cy_s.otype == t].kl.sum()), "Sales LY": float(ly_s[ly_s.otype == t].kl.sum())})
        t = pd.DataFrame(rows)
        tot_cy, tot_ly = t["Sales CY"].sum(), t["Sales LY"].sum()
        t["% of Total Sales CY"] = t["Sales CY"] / tot_cy if tot_cy else np.nan
        t["% of Total Sales LY"] = t["Sales LY"] / tot_ly if tot_ly else np.nan
        r["sec"] = t
        r["sec_label"] = (ws_sec["cy"].label if sec_period == "Month" else f"{mi_name(mi_of(fy, 1))} to {mi_name(mi)}") + \
                         (f" (till {sec_last:%d-%b-%y})" if ws_sec["cy"].partial else "")
        r["sec_other"] = float(cy_s[cy_s.otype == "Other"].kl.sum())
        out.append(r)
    return out


# ------------------------------------------------------------------ Excel in the same layout
_thin, _med = Side(style="thin"), Side(style="medium")
_B = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
_YEL = PatternFill("solid", fgColor="FFFF00")
_HEAD = PatternFill("solid", fgColor="DCE6F1")


def _sheet(ws, r: dict):
    for col, w in zip("ABCDEFGHI", (14, 11, 11, 11, 11, 11, 11, 12, 12)): ws.column_dimensions[col].width = w

    def cell(ref, v=None, bold=False, fill=None, fmt=None, align=_CENTER, size=11):
        c = ws[ref]; c.value = v
        c.font = Font(name="Calibri", size=size, bold=bold); c.alignment = align
        if fill: c.fill = fill
        if fmt: c.number_format = fmt
        return c

    def box(rng, border=_B):
        for row in ws[rng]:
            for c in row:
                c.border = border
                if c.value is None: c.alignment = _CENTER

    def merge(rng, v=None, **kw):
        ws.merge_cells(rng); cell(rng.split(":")[0], v, **kw); box(rng)

    merge("A2:I2", "*  2 - Sales Analysis", bold=True, size=12); ws["A2"].font = Font(name="Arial", size=12, bold=True); ws.row_dimensions[2].height = 21
    merge("A4:D4", "Name of the Stockist"); merge("E4:I4", r["stockist"], bold=True, fill=_YEL); ws.row_dimensions[4].height = 24.75
    for top, key, label in ((6, "month", r["month_label"] + (" (till " + f"{r['data_till']:%d-%b-%y}" + ")" if r["partial"] else "")),
                            (10, "cum", r["cum_label"])):
        if top == 6:
            merge("A6:C6", "Month", bold=True); merge("D6:I6", label, bold=True, fill=_YEL)
        else:
            merge("A10:C10", "Cumulative", bold=True); merge("D10:I10", label, bold=True, fill=_YEL)
        h = top + 1
        for col, txt in zip("ABCDE", ("Potential (KL)", f"Target (KL)", "Sales CY (KL)", "Sales LY (KL)", "(+/-)")): cell(f"{col}{h}", txt, fill=_HEAD)
        merge(f"F{h}:I{h}", "Remarks", fill=_HEAD); box(f"A{h}:E{h}"); ws.row_dimensions[h].height = 32
        d, v = r[key], h + 1
        for col, x in zip("ABCD", (d["potential"], d["target"], d["cy"], d["ly"])):
            cell(f"{col}{v}", None if x is None or x != x else round(float(x), 2), fmt="#,##0.00")
        cell(f"E{v}", f"=C{v}-D{v}", fmt="#,##0.00;[Red]-#,##0.00")
        merge(f"F{v}:I{v}", d["remark"], align=_LEFT, size=10); box(f"A{v}:E{v}"); ws.row_dimensions[v].height = 30
    ws.row_dimensions[13].height = 15.75
    merge("A14:I14", "LUBES RESELLER GROUP", bold=True)
    merge("A15:I15", "Analysis of Secondary Sales ")
    merge("A16:D16", "Period"); merge("E16:I16", r["sec_label"], bold=True, fill=_YEL)
    merge("A17:C17", "Channel", fill=_HEAD)
    for col, txt in zip("DEFGHI", ("Total No.", "Cum. No. Supplied", "Sales CY", "Sales LY", "% of Total Sales CY", "% of Total Sales LY")):
        cell(f"{col}17", txt, fill=_HEAD)
    box("A17:I17"); ws.row_dimensions[17].height = 31.5
    for i, row in enumerate(r["sec"].itertuples(index=False)):
        n = 18 + i
        merge(f"A{n}:C{n}", row.Channel, align=_LEFT)
        cell(f"D{n}", int(row[1])); cell(f"E{n}", int(row[2]))
        cell(f"F{n}", round(float(row[3]), 2), fmt="#,##0.00"); cell(f"G{n}", round(float(row[4]), 2), fmt="#,##0.00")
        cell(f"H{n}", f'=IF($F$24=0,"",F{n}/$F$24)', fmt="0.0%"); cell(f"I{n}", f'=IF($G$24=0,"",G{n}/$G$24)', fmt="0.0%")
        box(f"A{n}:I{n}")
    merge("A24:C24", "Total", bold=True)
    for col, f in zip("DEFG", ("=SUM(D18:D23)", "=SUM(E18:E23)", "=SUM(F18:F23)", "=SUM(G18:G23)")):
        cell(f"{col}24", f, bold=True, fmt="#,##0.00" if col in "FG" else "0")
    cell("H24", '=IF(F24=0,"",SUM(H18:H23))', bold=True, fmt="0.0%"); cell("I24", '=IF(G24=0,"",SUM(I18:I23))', bold=True, fmt="0.0%")
    box("A24:I24")
    # thick outline round the reseller block, as in the format
    for row in ws["A14:I24"]:
        for c in row:
            b = c.border
            c.border = Border(left=_med if c.column == 1 else b.left, right=_med if c.column == 9 else b.right,
                              top=_med if c.row == 14 else b.top, bottom=_med if c.row == 24 else b.bottom)
    note = (f"Sales in the top section = {'primary (company to stockist)' if r['basis'] == 'Primary' else 'secondary (stockist to counters)'}. "
            "Secondary excludes ClearBlue (DEF). Total No. = outlets billed in the 12 months to the period end "
            "(Secondary file). " + (f"Potential = {r['pot_fy']} sales {r['pot_base']:,.1f} KL / {r['share']:.0%} assumed IOCL market share "
            f"(range {r['share_range'][0]:.0%}-{r['share_range'][1]:.0%}: {r['pot_base'] / r['share_range'][1]:,.0f}-{r['pot_base'] / r['share_range'][0]:,.0f} KL a year), "
            "spread evenly over the months." if r["pot_src"] == "share" else "Potential from Data/SSR Potential.xlsx."))
    ws["A26"] = note; ws["A26"].font = Font(size=8, italic=True, color="666666"); ws["A26"].alignment = _LEFT
    ws.merge_cells("A26:I27")
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "portrait"; ws.page_setup.fitToWidth = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True


def _safe_title(s: str) -> str:
    return re.sub(r"[\[\]:*?/\\]", "", s)[:31] or "Sheet"


def workbook_bytes(blocks: list[dict]) -> bytes:
    wb = Workbook(); wb.remove(wb.active)
    for r in blocks: _sheet(wb.create_sheet(_safe_title(r["stockist"])), r)
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def file_name(r: dict) -> str:
    return f"LRS Sales Analysis - {re.sub(r'[^A-Za-z0-9 ]', '', r['stockist']).strip()} - {r['month_label'].replace(chr(39), '')}.xlsx"


def zip_bytes(blocks: list[dict]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for r in blocks: z.writestr(file_name(r), workbook_bytes([r]))
    return buf.getvalue()
